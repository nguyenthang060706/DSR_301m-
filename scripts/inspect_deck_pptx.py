import sys, io
from pptx import Presentation
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

prs = Presentation('outputs/Team_Meeting_Data_Collection_thu_thap_xu_ly_v2.pptx')
slide_width = prs.slide_width
slide_height = prs.slide_height
print(f'Width: {slide_width}, Height: {slide_height} ({slide_width/914400:.2f} x {slide_height/914400:.2f} inches)')

for i, slide in enumerate(prs.slides, 1):
    print(f'--- SLIDE {i} ---')
    for shape in slide.shapes:
        print(f'Shape {shape.shape_id} ({shape.name}, type={shape.shape_type}): left={shape.left/914400:.2f}, top={shape.top/914400:.2f}, width={shape.width/914400:.2f}, height={shape.height/914400:.2f}')
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    print(f'    P: "{txt}"')
